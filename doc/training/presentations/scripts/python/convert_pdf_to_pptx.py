#!/usr/bin/env python3
"""
Convert PDF presentations to PPTX format
Note: Each PDF page becomes an image slide in PPTX (not editable text)
"""

import sys
import os
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches
from io import BytesIO

def pdf_to_pptx(pdf_path, pptx_path):
    """Convert PDF to PPTX by embedding each page as an image"""
    print(f"Converting {pdf_path}...")

    try:
        # Convert PDF pages to images
        images = convert_from_path(pdf_path, dpi=150)

        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
        prs.slide_height = Inches(7.5)

        # Add each page as a slide
        for i, img in enumerate(images):
            print(f"  Processing page {i+1}/{len(images)}...")

            # Add blank slide
            blank_slide_layout = prs.slide_layouts[6]  # Blank layout
            slide = prs.slides.add_slide(blank_slide_layout)

            # Save image to BytesIO
            img_buffer = BytesIO()
            img.save(img_buffer, format='PNG')
            img_buffer.seek(0)

            # Add image to slide (full size)
            left = top = Inches(0)
            slide.shapes.add_picture(img_buffer, left, top,
                                    width=prs.slide_width,
                                    height=prs.slide_height)

        # Save presentation
        prs.save(pptx_path)
        print(f"  ✓ Created: {pptx_path}")
        return True

    except Exception as e:
        print(f"  ✗ Error: {str(e)}")
        return False

if __name__ == "__main__":
    # Get all course PDFs
    courses = [
        "Course_01_LoRa_LoRaWAN.pdf",
        "Course_02_LBM_Architecture.pdf",
        "Course_03_USP_RAC.pdf",
        "Course_04_Multiprotocol.pdf",
        "Course_05_Hardware_Integration.pdf",
        "Course_06_Advanced_Features.pdf"
    ]

    success_count = 0
    for pdf_file in courses:
        if os.path.exists(pdf_file):
            pptx_file = pdf_file.replace('.pdf', '.pptx')
            if pdf_to_pptx(pdf_file, pptx_file):
                success_count += 1
        else:
            print(f"Warning: {pdf_file} not found")

    print(f"\n✓ Converted {success_count}/{len(courses)} courses to PPTX")
    print("\nNote: PPTX files contain image slides (not editable text).")
    print("For fully editable slides, use the original .tex source files.")
